"""
EXPORTAR LISTAS DEL EXCEL A catalogos.json
==========================================
Lee del Excel de la empresa los hospitales (con sus sedes, dirección, teléfono,
ciudad y departamento), los municipios y las indicaciones de terapia, y los guarda
en catalogos.json junto a index.html. Así el formulario tiene las listas aunque
no esté corriendo el servidor (por ejemplo, en GitHub Pages).

Uso:
    python exportar_catalogos.py                      (busca la plantilla en esta carpeta)
    python exportar_catalogos.py "ruta\\al\\Excel.xlsx"

Lee las mismas hojas y columnas que el servidor (leer_catalogos en App.py).
"""

import json
import os
import sys

import openpyxl

CARPETA = os.path.dirname(os.path.abspath(__file__))
PLANTILLA = os.path.join(CARPETA, "IMPLANT REPORTS - PLANTILLA UNIANDES_0685.xlsx")
SALIDA = os.path.join(CARPETA, "catalogos.json")


def leer_catalogos(archivo):
    libro = openpyxl.load_workbook(archivo, data_only=True)

    # Municipio -> departamento (hoja Municipios: D = municipio, B = departamento)
    municipios = []
    depto_de = {}
    for fila in libro["Municipios"].iter_rows(min_row=2, values_only=True):
        if fila[3] and fila[1]:
            municipios.append([str(fila[3]).strip(), str(fila[1]).strip()])
            depto_de[str(fila[3]).strip()] = str(fila[1]).strip()

    # Nombre del hospital -> ciudad (hoja Cuentas: E = nombre, F = ciudad)
    ciudad_de = {}
    for fila in libro["Cuentas"].iter_rows(min_row=2, values_only=True):
        if fila[4] and fila[5]:
            ciudad_de[str(fila[4]).strip().upper()] = str(fila[5]).strip()

    # Hospitales y sedes (hoja Inst_: B = nombre, D = dirección, E = teléfono)
    hospitales = []
    for fila in libro["Inst_"].iter_rows(min_row=2, max_row=73, values_only=True):
        if fila[1]:
            nombre = str(fila[1]).strip()
            ciudad = ciudad_de.get(nombre.upper(), "")
            hospitales.append({"nombre": nombre,
                               "direccion": str(fila[3] or "").strip(),
                               "telefono": str(fila[4] or "").strip(),
                               "ciudad": ciudad,
                               "departamento": depto_de.get(ciudad, "")})

    # Indicaciones de terapia (hoja Implant Report, columna AF, filas 2 a 27)
    indicaciones = [str(c[0].value).strip() for c in libro["Implant Report"]["AF2:AF27"] if c[0].value]

    return {"hospitales": hospitales, "municipios": municipios, "indicaciones": indicaciones}


if __name__ == "__main__":
    archivo = sys.argv[1] if len(sys.argv) > 1 else PLANTILLA
    if not os.path.isfile(archivo):
        sys.exit("No encuentro el Excel: " + archivo)
    datos = leer_catalogos(archivo)
    with open(SALIDA, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)
    print(f"Listo: {len(datos['hospitales'])} hospitales/sedes, {len(datos['municipios'])} municipios, "
          f"{len(datos['indicaciones'])} indicaciones -> {SALIDA}")
