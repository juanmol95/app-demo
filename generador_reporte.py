#
"""
GENERADOR DEL IMPLANT REPORT
============================
"""

import json
import os
import re
import sys
import warnings
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# openpyxl avisa que ignora algunas validaciones del Excel. No afecta el resultado.
warnings.filterwarnings("ignore", category=UserWarning)


# CONFIGURACIÓN - lo que se puede cambiar fácilmente
CARPETA = os.path.dirname(os.path.abspath(__file__))        # la carpeta donde está este archivo
PLANTILLA = os.path.join(CARPETA, "IMPLANT REPORTS - PLANTILLA UNIANDES_0685.xlsx")   
CARPETA_SALIDA = os.path.join(CARPETA, "salidas")                # aquí se guardan los reportes
FILA = 3                                  # fila de la hoja "Datos" que lee el Implant Report
CONSERVAR_EXCEL = False                   # True = guardar también el Excel lleno junto al PDF


# Cada línea es: (columna del Excel, sección del JSON, campo del JSON, tipo de dato)
# Si la sección está vacía ("") el campo está en la raíz del JSON.
CAMPOS = [
    ("B",  "",            "especialista",     "texto"),
    ("E",  "paciente",    "diagnostico",      "texto"),
    ("F",  "paciente",    "fe",               "numero"),
    ("G",  "paciente",    "sexo",             "sexo"),
    ("H",  "paciente",    "fecha_nacimiento", "fecha"),
    ("I",  "paciente",    "documento",        "texto"),
    ("J",  "paciente",    "direccion",        "texto"),
    ("K",  "paciente",    "telefono",         "texto"),
    ("L",  "paciente",    "ciudad",           "texto"),
    ("M",  "paciente",    "departamento",     "texto"),
    ("O",  "medico",      "nombre",           "texto"),
    ("P",  "medico",      "apellido",         "texto"),
    ("Q",  "hospital",    "nombre",           "texto"),
    ("R",  "hospital",    "direccion",        "texto"),
    ("S",  "hospital",    "telefono",         "texto"),
    ("T",  "hospital",    "ciudad",           "texto"),
    ("U",  "hospital",    "departamento",     "texto"),
    ("W",  "dispositivo", "modelo",           "numero"),
    ("X",  "dispositivo", "serial",           "texto"),
    ("Y",  "dispositivo", "fabricante",       "texto"),
    ("AA", "dispositivo", "ubicacion",        "ubicacion"),
    ("AB", "parametros",  "modo",             "texto"),
    ("AC", "parametros",  "lrl",              "numero"),
    ("AD", "parametros",  "url",              "numero"),
    ("AE", "parametros",  "av_delay",         "numero"),
    ("AF", "parametros",  "pvarp",            "numero"),
    ("AG", "parametros",  "v_ref",            "numero"),
    ("AH", "parametros",  "av_search",        "onoff"),
    ("AI", "parametros",  "vrr",              "onoff"),
    ("AJ", "parametros",  "atr",              "onoff"),
    ("AK", "parametros",  "sbr",              "onoff"),
]

# Datos que la hoja "Datos" no tiene, pero el formulario oficial sí los pide.
# Se escriben directo en la hoja "Implant Report": (celda, sección del JSON, campo del JSON)
EXTRAS_REPORTE = [
    ("W7", "paciente", "sufijo"),            # Suffix del paciente
    ("AA9", "paciente", "medical_record"),   # Historia clínica 
    ]

# Los electrodos tienen 14 columnas seguidas cada uno.
ELECTRODOS = {
    "ventricular_derecho":   "AL",
    "auricular_derecho":     "AZ",
    "ventricular_izquierdo": "BN",
}
CAMPOS_ELECTRODO = [
    ("modelo", "numero"), ("serial", "texto"), ("fabricante", "texto"),
    ("fecha_implante", "fecha"), ("polaridad", "texto"), ("posicion", "texto"),
    ("onda", "numero"), ("paced", "numero"), ("impedancia", "numero"), ("hv", "numero"),
    ("ancho", "numero"), ("umbral", "numero"), ("dft", "numero"), ("afib", "texto"),
]

# Eexplantados: 6 columnas seguidas cada uno.
EXPLANTES = {
    "dispositivo":           "CB",
    "electrodo_auricular":   "CH",
    "electrodo_ventricular": "CN",
}
CAMPOS_EXPLANTE = [
    ("modelo", "numero"), ("serial", "texto"), ("fabricante", "texto"),
    ("razon", "texto"), ("estado", "texto"), ("fecha_implante", "fecha"),
]

# Datos que no pueden faltar. Si falta alguno, el programa avisa (pero igual genera el reporte).
# Lo demás (departamento, dirección del hospital...) se busca solo en las otras hojas del Excel.
OBLIGATORIOS = [
    ("", "especialista"), ("", "fecha_implante"),
    ("paciente", "documento"), ("paciente", "primer_nombre"),
    ("paciente", "primer_apellido"), ("paciente", "sexo"), ("paciente", "fecha_nacimiento"),
    ("medico", "nombre"), ("medico", "apellido"),
    ("hospital", "nombre"),
    ("dispositivo", "modelo"), ("dispositivo", "serial"),
]


# FUNCIONES PEQUEÑAS PARA ARREGLAR LOS DATOS
def sacar_dato(datos, seccion, campo):
    """Busca un campo dentro del JSON. Devuelve None si no existe."""
    if seccion == "":
        return datos.get(campo)
    return datos.get(seccion, {}).get(campo)


def fecha_a_numero(texto):
    """
    Convierte '19850314', '1985-03-14' o '14/03/1985' en el número que usa Excel
    para guardar fechas (los días que han pasado desde el 30/12/1899).
    """
    texto = str(texto).strip()
    for formato in ("%Y%m%d", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            fecha = datetime.strptime(texto, formato)
            return (fecha - datetime(1899, 12, 30)).days
        except ValueError:
            pass
    raise ValueError("No entiendo la fecha: " + texto)


def a_numero_si_se_puede(valor):
    """'540' -> 540,  '0.7' -> 0.7,  'G158' -> 'G158',  '0185' -> '0185' (queda texto)."""
    if isinstance(valor, (int, float)):
        return valor
    texto = str(valor).strip().replace(",", ".")
    sin_signo = texto.lstrip("-")
    es_numero = sin_signo.replace(".", "", 1).isdigit()
    tiene_cero_inicial = sin_signo.startswith("0") and len(sin_signo) > 1 and not sin_signo.startswith("0.")
    if es_numero and not tiene_cero_inicial:
        return float(texto) if "." in texto else int(texto)
    return str(valor).strip()


def arreglar_sexo(valor):
    """El Excel espera 'Male' o 'Female'."""
    valor = str(valor).strip().lower()
    if valor in ("m", "male", "masculino", "hombre"):
        return "Male"
    if valor in ("f", "female", "femenino", "mujer"):
        return "Female"
    return "Unknown"


def arreglar_ubicacion(valor):
    """El Excel espera, por ejemplo, 'Left  Subpectoral' (con dos espacios) o 'Right Subpectoral'."""
    partes = str(valor).strip().split(" ", 1)
    if len(partes) == 2 and partes[0].lower() in ("left", "izquierdo", "izq"):
        return "Left  " + partes[1].strip()
    if len(partes) == 2 and partes[0].lower() in ("right", "derecho", "der"):
        return "Right " + partes[1].strip()
    return str(valor).strip()


def convertir(valor, tipo):
    """Deja el dato listo para escribirlo en el Excel, según su tipo."""
    if tipo == "fecha":
        return fecha_a_numero(valor)
    if tipo == "numero":
        return a_numero_si_se_puede(valor)
    if tipo == "sexo":
        return arreglar_sexo(valor)
    if tipo == "ubicacion":
        return arreglar_ubicacion(valor)
    if tipo == "onoff":
        return "ON" if str(valor).strip().lower() in ("on", "si", "sí", "true", "1") else "OFF"
    return str(valor).strip()          # texto


def anotar(cambios, columna, valor, tipo):
    """
    Apunta un dato en el diccionario 'cambios' (columna -> valor).
    Todavía no escribe nada en el Excel: eso se hace al final, de una sola vez.
    Si el dato viene vacío, no hace nada.
    """
    if valor is None or str(valor).strip() == "":
        return
    cambios[columna] = convertir(valor, tipo)


# BUSCAR DATOS EN LAS OTRAS HOJAS DEL EXCEL (hospitales, ciudades, etc.)
def buscar(hoja, col_clave, col_resultado, clave):
    """
    Recorre una hoja fila por fila. Cuando en la columna 'col_clave' encuentra
    el valor 'clave', devuelve lo que hay en la columna 'col_resultado'.
    Si no lo encuentra devuelve None.
    """
    i_clave = column_index_from_string(col_clave) - 1
    i_resultado = column_index_from_string(col_resultado) - 1
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[i_clave] is not None and str(fila[i_clave]).strip().lower() == str(clave).strip().lower():
            return fila[i_resultado]
    return None


def buscar_datos_derivados(cambios, libro_valores):
    """
    Completa lo que NO se capturó, buscándolo en las hojas de apoyo:
      - departamento del paciente          (hoja Municipios)
      - dirección y teléfono del hospital  (hoja Inst_)
      - ciudad y departamento del hospital (hojas Cuentas y Municipios)
    Regla: si el dato ya viene capturado, se respeta. Las hojas solo rellenan lo que falta.
    """
    ciudad_paciente = cambios.get("L")
    hospital = cambios.get("Q")

    depto_paciente = buscar(libro_valores["Municipios"], "D", "B", ciudad_paciente) if ciudad_paciente else None
    direccion = buscar(libro_valores["Inst_"], "B", "D", hospital) if hospital else None
    telefono = buscar(libro_valores["Inst_"], "B", "E", hospital) if hospital else None
    ciudad_hospital = cambios.get("T") or (buscar(libro_valores["Cuentas"], "E", "F", hospital) if hospital else None)
    depto_hospital = buscar(libro_valores["Municipios"], "D", "B", ciudad_hospital) if ciudad_hospital else None

    encontrados = {"M": depto_paciente, "R": direccion, "S": telefono,
                   "T": ciudad_hospital, "U": depto_hospital}
    for columna, valor in encontrados.items():
        # Solo anotamos si lo encontramos y nadie lo había capturado
        if valor is not None and columna not in cambios:
            cambios[columna] = valor


# REVISAR LOS DATOS (avisos útiles antes de generar el reporte)
def revisar_datos(datos, libro_valores):
    """Devuelve una lista de avisos en español. Si la lista está vacía, todo está bien."""
    avisos = []

    for seccion, campo in OBLIGATORIOS:
        if not sacar_dato(datos, seccion, campo):
            avisos.append("Falta el dato: " + (seccion + " > " if seccion else "") + campo)

    hospital = sacar_dato(datos, "hospital", "nombre")
    if hospital and buscar(libro_valores["Inst_"], "B", "A", hospital) is None:
        avisos.append("El hospital '" + hospital + "' no está en la hoja Inst_ (revisa que esté bien escrito).")

    modelo = sacar_dato(datos, "dispositivo", "modelo")
    if modelo:
        modelos_conocidos = [str(fila[1]).strip() for fila in
                             libro_valores["PCA"].iter_rows(min_row=3, max_row=122, values_only=True)
                             if fila[1] is not None]
        if str(modelo).strip() not in modelos_conocidos:
            avisos.append("El modelo '" + str(modelo) + "' no está en la hoja PCA: no se sabrá el tipo de dispositivo.")

    # El Implant Report lee la fila (Consumo!A1 + 2) de Datos. Con A1 = 1 es la fila 3.
    if libro_valores["Consumo"]["A1"].value != 1:
        avisos.append("En la hoja Consumo, la celda A1 debe valer 1; si no, el reporte leerá otra fila.")

    return avisos


ARCHIVO_HOJA_DATOS = "xl/worksheets/sheet1.xml"     # aquí vive la hoja "Datos"
ARCHIVO_HOJA_REPORTE = "xl/worksheets/sheet3.xml"   # aquí vive la hoja "Implant Report"


def celda_xml(referencia, estilo, valor):
    """Arma el XML de una celda: un número, o un texto."""
    atributo_estilo = ' s="' + estilo + '"' if estilo else ""
    if isinstance(valor, (int, float)):
        return '<c r="' + referencia + '"' + atributo_estilo + '><v>' + repr(valor) + '</v></c>'
    # Quitamos caracteres raros que Excel no acepta en un texto
    texto = re.sub("[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(valor))
    return ('<c r="' + referencia + '"' + atributo_estilo + ' t="inlineStr"><is><t xml:space="preserve">'
            + escape(texto) + '</t></is></c>')


def celda_formula_xml(referencia, estilo, formula):
    """Arma el XML de una celda que contiene una fórmula, por ejemplo =T3."""
    atributo_estilo = ' s="' + estilo + '"' if estilo else ""
    return '<c r="' + referencia + '"' + atributo_estilo + ' t="str"><f>' + escape(formula) + '</f><v></v></c>'


def modificar_fila(xml, cambios, formulas):
    """Cambia las celdas de la fila FILA dentro del XML de la hoja."""
    fila = re.search(r'(<row r="' + str(FILA) + r'"[^>]*>)(.*?)(</row>)', xml, re.S)
    inicio_fila, contenido, fin_fila = fila.groups()

    # Separamos la fila en celdas, guardadas por número de columna (A=1, B=2...)
    celdas = {}
    for celda in re.findall(r"<c [^>]*?/>|<c [^>]*?>.*?</c>", contenido, re.S):
        columna = re.match(r'<c r="([A-Z]+)', celda).group(1)
        celdas[column_index_from_string(columna)] = celda

    def estilo_de(columna):
        """Cada celda tiene un 'estilo' (formato de fecha, tipo de letra...). Lo conservamos."""
        celda = celdas.get(column_index_from_string(columna), "")
        encontrado = re.search(r'\ss="(\d+)"', celda.split(">")[0])
        return encontrado.group(1) if encontrado else None

    for columna, valor in cambios.items():
        celdas[column_index_from_string(columna)] = celda_xml(columna + str(FILA), estilo_de(columna), valor)
    for columna, formula in formulas.items():
        celdas[column_index_from_string(columna)] = celda_formula_xml(columna + str(FILA), estilo_de(columna), formula)

    # Volvemos a armar la fila con las celdas en orden
    fila_nueva = inicio_fila + "".join(celdas[numero] for numero in sorted(celdas)) + fin_fila
    return xml[:fila.start()] + fila_nueva + xml[fila.end():]


def escribir_celdas_reporte(xml, extras):
    """
    Escribe un texto en celdas sueltas de la hoja "Implant Report" (por ejemplo W7).
    Esas celdas ya existen vacías en la plantilla, con su formato; solo les ponemos el texto.
    """
    for referencia, texto in extras.items():
        celda_vacia = re.search(r'<c r="' + referencia + r'"( s="\d+")?\s*/>', xml)
        if celda_vacia:
            estilo = re.search(r'\d+', celda_vacia.group(1) or "")
            nueva = celda_xml(referencia, estilo.group(0) if estilo else None, texto)
            xml = xml[:celda_vacia.start()] + nueva + xml[celda_vacia.end():]
    return xml


def guardar_excel(archivo_plantilla, archivo_nuevo, cambios, formulas, extras_reporte):
    """Copia la plantilla en 'archivo_nuevo' con la fila 3 de la hoja Datos modificada."""
    with zipfile.ZipFile(archivo_plantilla) as original, \
            zipfile.ZipFile(archivo_nuevo, "w", zipfile.ZIP_DEFLATED) as nuevo:
        for archivo in original.infolist():
            nombre = archivo.filename
            contenido = original.read(nombre)

            if nombre == "xl/calcChain.xml":
                # Lista interna de fórmulas. Como cambiamos algunas, la borramos
                # y Excel la vuelve a crear solo al abrir el archivo.
                continue
            if nombre == ARCHIVO_HOJA_DATOS:
                contenido = modificar_fila(contenido.decode("utf-8"), cambios, formulas).encode("utf-8")
            elif nombre == ARCHIVO_HOJA_REPORTE and extras_reporte:
                contenido = escribir_celdas_reporte(contenido.decode("utf-8"), extras_reporte).encode("utf-8")
            elif nombre == "xl/workbook.xml":
                # Le pedimos a Excel que recalcule todas las fórmulas al abrir
                texto = contenido.decode("utf-8")
                texto = re.sub(r"<calcPr[^>]*/>", '<calcPr calcId="191029" fullCalcOnLoad="1"/>', texto, count=1)
                contenido = texto.encode("utf-8")
            elif nombre == "xl/_rels/workbook.xml.rels":
                contenido = re.sub(r"<Relationship [^>]*calcChain[^>]*/>", "", contenido.decode("utf-8")).encode("utf-8")
            elif nombre == "[Content_Types].xml":
                contenido = re.sub(r"<Override [^>]*calcChain[^>]*/>", "", contenido.decode("utf-8")).encode("utf-8")

            nuevo.writestr(archivo, contenido)


# EXPORTAR A PDF CON EXCEL

def crear_pdf(archivo_excel, archivo_pdf):
    """Abre el Excel con la aplicación Excel y guarda la hoja 'Implant Report' como PDF."""
    import pythoncom                    # vienen con la librería pywin32
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False               # que no se vea la ventana
    excel.DisplayAlerts = False
    libro = None
    try:
        libro = excel.Workbooks.Open(os.path.abspath(archivo_excel), UpdateLinks=0, ReadOnly=True)
        excel.CalculateFull()           # recalcula todas las fórmulas
        libro.Worksheets("Implant Report").ExportAsFixedFormat(0, os.path.abspath(archivo_pdf))  # 0 = PDF
    finally:
        if libro is not None:
            libro.Close(False)
        excel.Quit()


# PROGRAMA PRINCIPAL
def preparar_cambios(datos):
    """Recorre el JSON y arma el diccionario {columna: valor} con todo lo que hay que escribir."""
    cambios = {}

    # Los campos sueltos (paciente, médico, hospital, dispositivo, parámetros)
    for columna, seccion, campo, tipo in CAMPOS:
        anotar(cambios, columna, sacar_dato(datos, seccion, campo), tipo)

    # El nombre y el apellido del paciente: el Excel los usa en una sola celda
    paciente = datos.get("paciente", {})
    nombre = (paciente.get("primer_nombre", "") + " " + paciente.get("segundo_nombre", "")).strip()
    apellido = (paciente.get("primer_apellido", "") + " " + paciente.get("segundo_apellido", "")).strip()
    anotar(cambios, "C", nombre.upper(), "texto")
    anotar(cambios, "D", apellido.upper(), "texto")

    # La fecha del procedimiento (columna A) y la fecha de implante del dispositivo (Z)
    anotar(cambios, "A", datos.get("fecha_implante"), "fecha")
    fecha_dispositivo = sacar_dato(datos, "dispositivo", "fecha_implante") or datos.get("fecha_implante")
    anotar(cambios, "Z", fecha_dispositivo, "fecha")

    # Si no dicen el fabricante del dispositivo, asumimos Boston Scientific
    if sacar_dato(datos, "dispositivo", "modelo") and not sacar_dato(datos, "dispositivo", "fabricante"):
        anotar(cambios, "Y", "Boston Scientific", "texto")

    # Electrodos y explantes: recorremos cada bloque de columnas seguidas
    for bloques, lista_campos, nombre_seccion in ((ELECTRODOS, CAMPOS_ELECTRODO, "electrodos"),
                                                  (EXPLANTES, CAMPOS_EXPLANTE, "explantes")):
        for nombre_bloque, primera_columna in bloques.items():
            numero_inicial = column_index_from_string(primera_columna)
            for posicion, (campo, tipo) in enumerate(lista_campos):
                columna = get_column_letter(numero_inicial + posicion)
                valor = datos.get(nombre_seccion, {}).get(nombre_bloque, {}).get(campo)
                anotar(cambios, columna, valor, tipo)

    return cambios


def generar_reporte(datos):
    """Recibe los datos (un diccionario de Python) y genera el PDF (y opcionalmente el Excel)."""
    os.makedirs(CARPETA_SALIDA, exist_ok=True)

    # Abrimos la plantilla solo para LEER valores (hospitales, ciudades, modelos...)
    libro_valores = openpyxl.load_workbook(PLANTILLA, data_only=True)

    avisos = revisar_datos(datos, libro_valores)
    cambios = preparar_cambios(datos)
    buscar_datos_derivados(cambios, libro_valores)

    # Corrección: en la plantilla, "Datos del Hospital" (columnas DX, DY, DZ) copia la
    # ciudad del PACIENTE (L, M, N). Aquí las apuntamos a la ciudad del HOSPITAL (T, U, V).
    formulas = {"DX": "T" + str(FILA), "DY": "U" + str(FILA), "DZ": "V" + str(FILA)}

    # Guardamos la copia con un nombre único (cédula + fecha y hora)
    documento = re.sub(r"\W+", "", str(sacar_dato(datos, "paciente", "documento") or "sin_documento"))
    nombre_base = "Implant_Report_" + documento + "_" + datetime.now().strftime("%Y%m%d_%H%M%S")
    archivo_excel = os.path.join(CARPETA_SALIDA, nombre_base + ".xlsx")
    archivo_pdf = os.path.join(CARPETA_SALIDA, nombre_base + ".pdf")

    # Datos que van directo en la hoja "Implant Report" (sufijo, historia clínica)
    extras_reporte = {}
    for celda, seccion, campo in EXTRAS_REPORTE:
        valor = sacar_dato(datos, seccion, campo)
        if valor is not None and str(valor).strip() != "":
            extras_reporte[celda] = str(valor).strip()

    guardar_excel(PLANTILLA, archivo_excel, cambios, formulas, extras_reporte)

    # Exportamos a PDF. Si falla, nos quedamos con el Excel.
    pdf_creado = False
    try:
        crear_pdf(archivo_excel, archivo_pdf)
        pdf_creado = True
        if not CONSERVAR_EXCEL:
            os.remove(archivo_excel)
    except Exception as error:
        avisos.append("No se pudo crear el PDF (" + str(error) + "). Se dejó el Excel para exportarlo a mano.")

    return {
        "pdf": archivo_pdf if pdf_creado else None,
        "excel": archivo_excel if (CONSERVAR_EXCEL or not pdf_creado) else None,
        "avisos": avisos,
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: py generador_reporte_simple.py ejemplo_reporte.json")
        sys.exit(1)

    with open(sys.argv[1], encoding="utf-8") as archivo:
        datos_paciente = json.load(archivo)

    resultado = generar_reporte(datos_paciente)
    print(json.dumps(resultado, indent=2, ensure_ascii=False))
